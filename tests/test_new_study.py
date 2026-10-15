"""New study asks for the marker names and the sampling workbook up front.

Before, New study asked for neither: the markers stayed as the dye names (or a
template placeholder like Marker1) until someone found Set up > Channels and
markers, and the sampling workbook was first looked at by 06b at Unblind, the
very end of a study. Now both are on the dialog, the markers filled in from the
slide headers once read, and Create saves them where Set up keeps them.

    python tests/test_new_study.py
"""

import os
import sys
import tempfile
import time

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
sys.path[:0] = [os.path.join(REPO, "app"), os.path.join(REPO, "scripts"), HERE]
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from _fixture import use_temp_study                             # noqa: E402

STUDY = use_temp_study()
os.environ["LS_HOME"] = tempfile.mkdtemp()

failures = []


def chk(name, got, want=True):
    ok = got == want
    print(f"{'ok  ' if ok else 'FAIL'} {name:<70} {got!r}" + ("" if ok else f"  (wanted {want!r})"))
    if not ok:
        failures.append(name)


# Reuse the header writers from test_autodetect (everything above its first scenario).
src = open(os.path.join(HERE, "test_autodetect.py")).read().split('print("--- one scan per marker')[0]
ns = {"__name__": "autodetect_helpers", "__file__": os.path.join(HERE, "test_autodetect.py")}
exec(compile(src, "test_autodetect_helpers", "exec"), ns)
xml, write_czi = ns["xml"], ns["write_czi"]
DAPI, AF568, AF488 = ns["DAPI"], ns["AF568"], ns["AF488"]

import autodetect as AD                                         # noqa: E402

PAIRED = os.path.join(STUDY.root, "paired")
MULTI = os.path.join(STUDY.root, "multi")
os.makedirs(PAIRED)
os.makedirs(MULTI)
for n in range(1, 3):
    write_czi(os.path.join(PAIRED, "LS22_%da.czi" % n), xml([DAPI, AF568]))
    write_czi(os.path.join(PAIRED, "LS22_%db.czi" % n), xml([DAPI, AF488]))
    write_czi(os.path.join(MULTI, "LS3%d_1a.czi" % n), xml([DAPI, AF488, AF568]))

print("--- the markers the headers name ---")
chk("paired: one per scan's marker channel", AD.markers_found(AD.detect(PAIRED)),
    ["AF568", "AF488"])
chk("multiplex: the marker channels of the table", AD.markers_found(AD.detect(MULTI)),
    ["AF488", "AF568"])
chk("nothing read: none", AD.markers_found(None), [])

print("--- naming them on a new project ---")
p, found = AD.new_project(PAIRED, os.path.join(STUDY.root, "o1"), 14, detection=AD.detect(PAIRED))
chk("paired names nothing wrong", AD.name_markers(p, {"AF568": "pERK", "AF488": "PCNA"}), [])
# Paired scans are matched to their marker by the channel name, so the dye
# name stays the marker and the antibody is what it stains - as the channel
# editor writes it.
chk("...the markers keep the scan names pairing needs", p.get("acquisition.markers"),
    ["AF568", "AF488"])
chk("...and what each stains is recorded", p.get("marker_identity"),
    {"AF568": "pERK", "AF488": "PCNA"})
chk("...the study is still valid", p.validate()[0], [])

p, found = AD.new_project(MULTI, os.path.join(STUDY.root, "o2"), 14, detection=AD.detect(MULTI))
chk("multiplex names nothing wrong", AD.name_markers(p, {"AF488": "PCNA", "AF568": "pERK"}), [])
chan = {c["czi_name"]: c["name"] for c in p.get("acquisition.channels")}
chk("...the channels are renamed, still found by their dye name", chan,
    {"DAPI": "DAPI", "AF488": "PCNA", "AF568": "pERK"})
chk("...the study is still valid", p.validate()[0], [])
p2, _ = AD.new_project(MULTI, os.path.join(STUDY.root, "o3"), 14, detection=AD.detect(MULTI))
chk("left as found changes nothing",
    (AD.name_markers(p2, {"AF488": "AF488", "AF568": ""}),
     [c["name"] for c in p2.get("acquisition.channels")]), ([], ["DAPI", "AF488", "AF568"]))
chk("two markers given one name is refused",
    bool(AD.name_markers(p2, {"AF488": "pERK", "AF568": "pERK"})), True)

print("--- the dialog ---")
from PySide6.QtCore import QCoreApplication                     # noqa: E402
from PySide6.QtWidgets import QApplication, QMessageBox         # noqa: E402

app = QApplication.instance() or QApplication([])
warned = []
QMessageBox.warning = staticmethod(lambda p, t, x, *a, **k: warned.append(x))
QMessageBox.question = staticmethod(lambda *a, **k: QMessageBox.Yes)

from new_study import NewStudyDialog                            # noqa: E402


def settle(dlg):
    end = time.monotonic() + 20
    while dlg.slide_read.busy and time.monotonic() < end:
        app.processEvents()
        time.sleep(0.005)
    for _ in range(5):
        QCoreApplication.processEvents()


book = os.path.join(STUDY.root, "sampling.xlsx")
open(book, "wb").close()
dlg = NewStudyDialog()
chk("no markers shown before the slides are read", dlg.marker_names(), {})
dlg.slides.edit.setText(PAIRED)
dlg._read_slides()
settle(dlg)
chk("after the read, a name per marker, the dye name to start with", dlg.marker_names(),
    {"AF568": "AF568", "AF488": "AF488"})
dlg.marker_edits["AF568"].setText("pERK")
chk("the dialog has a sampling workbook field", hasattr(dlg, "workbook"), True)
dlg.workbook.setValue(book)
out = os.path.join(STUDY.root, "made")
dlg.output.edit.setText(out)
dlg.thickness.setValue(14.0)
dlg._create()
chk("the study is made", os.path.isfile(AD.study_file(out)), True)
proj = dlg.project
chk("...with the marker named", (proj.get("marker_identity") or {}).get("AF568"), "pERK")
chk("...the other left as it was", "AF488" in (proj.get("marker_identity") or {}), False)
chk("...and the workbook kept as 06b's own option",
    proj.stage_options("join_sampling"), {"--xlsx": book})

import ls_project as LP                                         # noqa: E402
again = LP.Project.open(AD.study_file(out))
chk("both are in the saved study file",
    (again.get("marker_identity"), again.stage_options("join_sampling")),
    ({"AF568": "pERK"}, {"--xlsx": book}))

dlg = NewStudyDialog()
dlg.slides.edit.setText(MULTI)
dlg._read_slides()
settle(dlg)
dlg.marker_edits["AF488"].setText("pERK")
dlg.marker_edits["AF568"].setText("pERK")
out2 = os.path.join(STUDY.root, "made2")
dlg.output.edit.setText(out2)
dlg.thickness.setValue(14.0)
warned.clear()
dlg._create()
chk("one name for two markers is not made", os.path.exists(AD.study_file(out2)), False)
chk("...and says why", any("pERK" in w for w in warned), True)
chk("...with no workbook, none is kept", dlg.workbook.value(), "")

print("--- Set up > Slides: Check workbook ---")
# 06b's matching run early, on the fish the study has, writing nothing.
import openpyxl                                                 # noqa: E402
import workbook_check as WC                                     # noqa: E402

LIVE = ["Fish ID", "timepoint", "enviroment", "brackish Tank", "sea water  tank", "treatment",
        "Body weight(g)", "Fork Length (cm)", "Sex", "Brain for", "slicing IHC July 2025"]
good = os.path.join(STUDY.root, "groups.xlsx")
wb = openpyxl.Workbook()
wb.active.append(LIVE)
wb.active.append([22, 1, "brackish", "305", None, "control", None, None, "f", "IHC", "x"])
wb.save(good)

import main_window as MW                                        # noqa: E402

w = MW.MainWindow(REPO, config_path=STUDY.path)
s = w.setup
project = s.session.project
project.set("source_dir", PAIRED, by="test")
for key in ("slide_naming.pattern", "slide_naming.case_insensitive"):
    project.set(key, AD.detect(PAIRED).values[key], by="test")
chk("the study's fish, from the slide names before any manifest",
    WC.study_animals(project.cfg), ({"LS22": 4}, "slide file names"))
chk("Set up shows the workbook next to a Check button",
    hasattr(s, "workbook_field") and hasattr(s, "check_workbook_btn"), True)
s.workbook_field.setValue(good)
s.workbook_field.valueChanged.emit(good)
chk("...and keeps it as 06b's own option", project.stage_options("join_sampling"),
    {"--xlsx": good})
before = sorted(os.listdir(STUDY.root))
s.check_workbook_btn.click()
chk("...the check runs off the window's thread", s.workbook_check.busy, True)
end = time.monotonic() + 20
while s.workbook_check.busy and time.monotonic() < end:
    app.processEvents()
    time.sleep(0.005)
app.processEvents()
chk("...and says how many were matched", "1 of 1 fish (4 of 4 slides) matched to a group"
    in s.workbook_line.text(), True)
chk("...writing nothing", sorted(os.listdir(STUDY.root)), before)
chk("...not in the results either",
    os.path.exists(os.path.join(project.get("out_root"), "results", "animal_metadata.csv")),
    False)
s.workbook_field.setValue(os.path.join(STUDY.root, "none.xlsx"))
s.check_workbook_btn.click()
end = time.monotonic() + 20
while s.workbook_check.busy and time.monotonic() < end:
    app.processEvents()
    time.sleep(0.005)
app.processEvents()
chk("a workbook that is not there is said so", "no workbook" in s.workbook_line.text(), True)

print("\n%s" % ("FAILED: %s" % failures if failures else "all passed"))
sys.exit(1 if failures else 0)
